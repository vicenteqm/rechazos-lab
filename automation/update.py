"""Rechazos Lab: lectura de cinco carpetas, cruce validado y publicación SOLO local."""
from __future__ import annotations
import argparse, calendar, csv, hashlib, json, os, re, shutil, subprocess, sys, tempfile, time, unicodedata
from collections import Counter, defaultdict
from datetime import datetime, date, timezone, timedelta
from pathlib import Path
import openpyxl

PROJECT = Path(__file__).resolve().parents[1]
PRIVATE = PROJECT / '.private'
CONFIG = PROJECT / 'automation' / 'config.json'
MONTHS = ['ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO','SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE']
VERSION = 1

def now():
    return datetime.now(timezone(timedelta(hours=-6)))

def norm(v):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFD',str(v or '')) if not unicodedata.combining(c)).upper().split())

def text(v):
    if v is None: return ''
    if isinstance(v,(datetime,date)): return v.isoformat()[:10]
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def key(v):
    return re.sub(r'\.0+$','',text(v).lstrip("'"))

def payroll(v):
    s=key(v).lstrip('*').strip()
    return s.lstrip('0') or '0' if s.isdigit() else ''

def region(v):
    return norm(v).replace('REGION ','') or 'Sin región en base'

def iso(v):
    if isinstance(v,(datetime,date)): return v.isoformat()[:10]
    if not v: return ''
    s=text(v)
    for fmt in ('%Y-%m-%d','%d/%m/%Y','%d-%m-%Y'):
        try: return datetime.strptime(s[:10],fmt).date().isoformat()
        except ValueError: pass
    return ''

def write_json(path,data):
    path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_suffix(path.suffix+'.tmp')
    tmp.write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
    os.replace(tmp,path)

def digest(path):
    h=hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''): h.update(chunk)
    return h.hexdigest()

def period(path):
    s=norm(path.stem)
    y=re.findall(r'(?:^|\D)(20\d{2})(?:\D|$)',str(path))
    year=y[-1] if y else str(now().year)
    compact=re.search(r'(20\d{2})(0[1-9]|1[0-2])[0-3]\d',s)
    if compact: return f'{compact[1]}-{compact[2]}'
    numeric=re.search(r'(20\d{2})[-_](0[1-9]|1[0-2])(?:[-_]|$)',s)
    if numeric: return f'{numeric[1]}-{numeric[2]}'
    for m,name in enumerate(MONTHS,1):
        if name in s: return f'{year}-{m:02d}'
    # ASISTENCIA_INVEX_10, ASISTENCIA_INVEX_21_10_2026 and numeric month files.
    match=re.search(r'(?:^|[_ -])(\d{1,2})[_ -](\d{1,2})[_ -](20\d{2})(?:\D|$)',s)
    if match:
        return f'{match[3]}-{int(match[2]):02d}' if 1<=int(match[2])<=12 else ''
    match=re.search(r'(?:ASISTENCIA(?:_INVEX)?|BPI|RECORDS?)[_ -](0?[1-9]|1[0-2])(?:$|[_ -])',s)
    return f'{year}-{int(match[1]):02d}' if match else ''

def discover(config):
    output={}
    for role,folder in config['folders'].items():
        directory=Path(folder)
        if not directory.is_dir(): raise ValueError(f'No existe la carpeta de {role}: {directory}')
        files=sorted(p for p in directory.rglob('*') if p.is_file() and p.suffix.lower() in ('.xlsx','.xlsm','.csv') and not p.name.startswith('~$'))
        if not files: raise ValueError(f'No hay archivos compatibles en la carpeta de {role}.')
        if role in ('stores','agency'):
            if len(files)!=1: raise ValueError(f'{role}: se esperaba un único archivo, hay {len(files)}. Retira las copias anteriores de esa carpeta.')
            output[role]=files
        else:
            groups=defaultdict(list)
            for p in files:
                month=period(p)
                if not month: raise ValueError(f'No se puede identificar el mes de {p.name}. Usa un nombre con mes y una carpeta con año.')
                groups[month].append(p)
            # Each month is a snapshot, never append an older snapshot to the new one.
            output[role]=[max(paths,key=lambda p:(p.stat().st_mtime_ns,p.name)) for month,paths in sorted(groups.items())]
    return output

def excel_tables(path):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    if wb.epoch.year==1904:
        wb.close();raise ValueError(f'{path.name}: calendario 1904 no admitido. Convierte al calendario 1900.')
    return wb

def table_rows(sheet,header_marker):
    # Some source files format the last Excel row. Stop after the actual table.
    sheet.reset_dimensions()
    header=None;blanks=0
    for values in sheet.iter_rows(values_only=True):
        if header is None:
            if header_marker([norm(v) for v in values]): header=list(values);yield header
            continue
        if not any(v is not None and text(v) for v in values):
            blanks+=1
            if blanks>=256: break
            continue
        blanks=0;yield list(values)

def headers(values):
    h=defaultdict(list)
    for i,v in enumerate(values): h[norm(v)].append(i)
    return h

def get(row,h,*names,last=False):
    for name in names:
        if norm(name) in h:
            idx=h[norm(name)][-1 if last else 0]
            return row[idx] if idx<len(row) else None
    return None

def required(h,names,path):
    missing=[n for n in names if norm(n) not in h]
    if missing: raise ValueError(f'{path.name}: faltan columnas {missing}')

ALIASES={
 'bpi':{'folio':['FOLIO','FOLIO FICO'],'created':['FECHA DE CREACION','F. CREACION'],'region':['REGION'],'store':['SEDE','TIENDA'],'det':['DETERMINANTE','DETE'],'status':['ESTATUS BANCO'],'employeeId':['ID'],'employeeUser':['USUARIO FICO'],'employeeAlias':['EJ USUARIO'],'employeeExecutive':['EJECUTIVO P']},
 'records':{'folio':['FOLIO FICO','FOLIO'],'created':['F. CREACION','FECHA DE CREACION'],'status':['ESTATUS FICO'],'approval':['APROBACION'],'process':['ESTATUS DEL PROCESO']},
 'agency':{'folio':['FOLIO','FOLIO FICO'],'reason':['ESTATUS_P','RAZON AGRUPADA','MOTIVO DE RECHAZO'],'status':['ESTATUS'],'fico':['ESTATUS_FICO'],'cut':['ARCHIVO_ULTIMO_CORTE','FECHA DE CORTE']}}
MANDATORY={'bpi':['folio','created','region','det','status'],'records':['folio','status'],'agency':['folio','reason']}

def read_standard(path,role):
    fields=ALIASES[role]
    wb=None
    try:
        if path.suffix.lower()=='.csv':
            raw=path.read_bytes()
            try: content=raw.decode('utf-8-sig')
            except UnicodeDecodeError: content=raw.decode('cp1252')
            delim=';' if content.splitlines()[0].count(';')>content.splitlines()[0].count(',') else ','
            rows=list(csv.reader(content.splitlines(),delimiter=delim)); candidates=[('CSV',iter(rows))]
        else:
            wb=excel_tables(path)
            candidates=[]
            for sh in wb.worksheets:
                preview=list(sh.iter_rows(min_row=1,max_row=12,values_only=True))
                score=max((sum(any(alias in [norm(v) for v in row] for alias in fields[f]) for f in MANDATORY[role]) for row in preview),default=0)
                if score==len(MANDATORY[role]): candidates.append((sh.title,table_rows(sh,lambda hs:any(x in hs for x in fields['folio']))))
        if not candidates: raise ValueError(f'{path.name}: no hay una hoja con las columnas de {role}.')
        if len(candidates)>1:
            matching=[c for c in candidates if period(path) and norm(c[0]) in (MONTHS[int(period(path)[5:])-1],)]
            if len(matching)==1: candidates=matching
            else: raise ValueError(f'{path.name}: varias hojas válidas; no se eligió una arbitrariamente.')
        name,iterator=candidates[0];head=next(iterator,None)
        if not head: raise ValueError(f'{path.name}: tabla vacía.')
        h=headers(head);mapping={f:next((a for a in aliases if a in h),'') for f,aliases in fields.items()}
        for f in MANDATORY[role]:
            if not mapping[f]: raise ValueError(f'{path.name}: falta la columna {f}')
        result=[{f:text(get(row,h,mapping[f])) if mapping[f] else '' for f in fields} for row in iterator]
        return {'role':role,'name':path.name,'sheet':name,'map':{f:f for f in fields},'rows':result,'period':period(path)}
    finally:
        if wb: wb.close()

def read_stores(path):
    wb=excel_tables(path)
    try:
        for sh in wb.worksheets:
            it=table_rows(sh,lambda hs:'NUMERO DE TIENDA/DETERMINANTE' in hs and 'EJECUTIVA RH' in hs)
            head=next(it,None)
            if head is None: continue
            h=headers(head);required(h,['REGION','TIENDA','TEAM LIDER','LIDER SWAT','EJECUTIVA RH','EJECUTIVA CPRO','EJECUTIVA ENLACE BANCO'],path)
            if len(h['SUBLIDER'])<2: raise ValueError('Base de tiendas: se requieren la columna numérica y la columna de nombre de SUBLIDER.')
            rows=[]
            for r in it:
                if norm(get(r,h,'PROYECTO')) not in ('','INVEX'): continue
                det=key(get(r,h,'NUMERO DE TIENDA/DETERMINANTE'))
                if not det.isdigit(): continue
                row={'det':det,'store':norm(str(get(r,h,'SOCIO') or '')+' '+str(get(r,h,'TIENDA') or '')),'region':region(get(r,h,'REGION'))}
                for field,col in [('teamLeader','TEAM LIDER'),('swatLeader','LIDER SWAT'),('subleader','SUBLIDER'),('rh','EJECUTIVA RH'),('cpro','EJECUTIVA CPRO'),('bankLink','EJECUTIVA ENLACE BANCO')]:
                    value=norm(get(r,h,col,last=field=='subleader'))
                    row[field]=value if value not in ('','PENDIENTE','N/A','SIN ASIGNAR','SIN SUBLIDER','0') else 'Sin asignar en base'
                rows.append(row)
            return {'role':'stores','name':path.name,'sheet':sh.title,'rows':rows}
        raise ValueError(f'{path.name}: no hay una hoja de base de tiendas reconocida.')
    finally: wb.close()

def read_attendance(path):
    wb=excel_tables(path)
    try:
        for sh in wb.worksheets:
            it=table_rows(sh,lambda hs:'NOMBRE COMPLETO' in hs and 'DETERMINANTE' in hs)
            head=next(it,None)
            if head is None: continue
            h=headers(head);required(h,['NOMBRE COMPLETO','DETERMINANTE','USUARIO','NUM DE NOMINA','PERFIL DE PUESTO BANCO'],path)
            days={str(d):h[str(d)][0] for d in range(1,32) if str(d) in h};rows=[]
            for r in it:
                name=norm(get(r,h,'NOMBRE COMPLETO'));det=key(get(r,h,'DETERMINANTE'))
                if not name or not det.isdigit(): continue
                rows.append({'name':name,'det':det,'role':norm(get(r,h,'PERFIL DE PUESTO BANCO')),'id':payroll(get(r,h,'NUM DE NOMINA')),'user':norm(get(r,h,'USUARIO')),'alias':norm(get(r,h,'USUARIO BPI')),'start':iso(get(r,h,'FECHA DE INGRESO A PDV')),'end':iso(get(r,h,'FECHA DE BAJA')),'status':norm(get(r,h,'ESTATUS')),'days':{d:norm(r[i]) for d,i in days.items() if i<len(r) and norm(r[i]) not in ('','0')}})
            return {'role':'attendance','name':path.name,'sheet':sh.title,'period':period(path),'rows':rows}
        raise ValueError(f'{path.name}: asistencia sin encabezados reconocidos.')
    finally: wb.close()

def read_cached(path,role,sha):
    cache=PRIVATE/'cache'/f'{VERSION}-{role}-{sha}.json'
    if cache.exists(): return json.loads(cache.read_text(encoding='utf-8'))
    stat=path.stat()
    doc=read_stores(path) if role=='stores' else read_attendance(path) if role=='attendance' else read_standard(path,role)
    after=path.stat()
    if (stat.st_size,stat.st_mtime_ns)!=(after.st_size,after.st_mtime_ns): raise ValueError(f'{path.name}: cambió durante la lectura; se conservará el tablero anterior.')
    write_json(cache,doc);return doc

def enrich(model,stores,attendance,bpi_docs):
    catalog=defaultdict(list)
    for s in stores['rows']: catalog[s['det']].append(s)
    # Ambiguous determinants must never multiply folios or choose the first store.
    storemap={}
    for det,values in catalog.items():
        unique={json.dumps(s,sort_keys=True):s for s in values}
        if len(unique)==1: storemap[det]=next(iter(unique.values()))
    ix=defaultdict(list);subix=defaultdict(list)
    for doc in attendance:
        for e in doc['rows']:
            e={**e,'source':doc['name']}
            for kind,value in [('id',e['id']),('user',e['user']),('alias',e['alias']),('name',e['name'])]:
                if value and value not in ('PENDIENTE','SIN USUARIO','N/A','0'): ix[(doc['period'],e['det'],kind,value)].append(e)
            if 'SUBLIDER' in e['role']: subix[(doc['period'],e['det'])].append(e)
    creators={}
    for d in bpi_docs:
        for r in d['rows']:
            if key(r['folio']).isdigit():
                previous=creators.get(key(r['folio']))
                if previous and any(previous.get(f)!=r.get(f) for f in ('employeeId','employeeUser','employeeAlias','employeeExecutive')):
                    raise ValueError('Identificadores del capturista distintos para el mismo folio BPI.')
                creators[key(r['folio'])]=r
    totals=Counter();by_month=defaultdict(Counter)
    def valid(e,day): return (not e['start'] or e['start']<=day) and (not e['end'] or e['end']>=day)
    def identity(e): return e['name']
    for row in model['rows']:
        det=row['det'];s=storemap.get(det);row['regionBpi']=row['region'];row['storeBpi']=row['store']
        for f in ('teamLeader','swatLeader','subleader','rh','cpro','bankLink'): row[f]=s[f] if s else 'Sin coincidencia en base'
        row['storeMatch']='Coincidencia única' if s else 'Determinante ambiguo' if det in catalog else 'Sin coincidencia'
        if s: row.update(region=s['region'],store=s['store']);totals['storesMatched']+=1
        else: totals['storesUnmatched']+=1
        r=creators[row['folio']];candidates=[];method=''
        matches=[]
        for kind,value in [('id',payroll(r.get('employeeId'))),('user',norm(r.get('employeeUser'))),('alias',norm(r.get('employeeAlias'))),('alias',norm(r.get('employeeExecutive'))),('name',norm(r.get('employeeExecutive')))]:
            found=[e for e in ix.get((row['month'],det,kind,value),[]) if valid(e,row['created'])] if value else []
            if found: matches.append((kind,found))
        if matches:
            identities=[{identity(e) for e in found} for kind,found in matches]
            common=set.intersection(*identities)
            if len(common)==1:
                person=next(iter(common));candidates=[e for kind,found in matches for e in found if identity(e)==person];method=' + '.join(dict.fromkeys(kind for kind,_ in matches))
        fallback=norm(r.get('employeeExecutive') or r.get('employeeAlias') or r.get('employeeUser')) or 'Sin identificador en BPI'
        if candidates:
            roles={e['role'] for e in candidates if e['role']};codes={e['days'].get(str(int(row['created'][-2:])),'') for e in candidates}-{''}
            row.update(promoter=candidates[0]['name'],employeeRole=next(iter(roles)) if len(roles)==1 else 'Puesto ambiguo',attendanceMatch='Coincidencia exacta',attendanceCode=next(iter(codes)) if len(codes)==1 else 'Sin captura' if not codes else 'Capturas distintas',attendanceSource=' | '.join(sorted({e['source'] for e in candidates})),employeeMatchMethod=method)
            totals['employeesMatched']+=1;by_month[row['month']]['employeesMatched']+=1
        else:
            row.update(promoter=f'{fallback} [sin validar]',employeeRole='Sin validar en asistencia',attendanceMatch='Identidad ambigua' if matches else 'Sin coincidencia exacta',attendanceCode='Sin validar',attendanceSource='',employeeMatchMethod='')
            totals['employeesUnmatched']+=1;by_month[row['month']]['employeesUnmatched']+=1
        if row['subleader']=='Sin asignar en base':
            subs={e['name'] for e in subix.get((row['month'],det),[]) if valid(e,row['created']) and e['days'].get(str(int(row['created'][-2:]))) in ('A','R')}
            if len(subs)==1: row['subleader']=next(iter(subs));row['subleaderSource']='Asistencia del día'
        row.setdefault('subleaderSource','Base vigente de tiendas')
    model['enrichment']={'totals':dict(totals),'byMonth':{m:dict(c) for m,c in by_month.items()},'ambiguousStores':sorted(set(catalog)-set(storemap)),'method':'Región y responsables de base vigente; persona por determinante + mes + identificadores exactos de BPI y asistencia. Sin coincidencias aproximadas. No se usa el nombre del cliente.'}
    return model

def sync_site(output):
    output.mkdir(parents=True,exist_ok=True)
    for name in ['index.html','styles.css','halloween.css','visuals.css','excel.css','pumpkin.svg','app.js','core.js','folder-ui.js','html-export.js','excel-ui.js','excel-worker.js','excel-export.js','README.md']:
        shutil.copy2(PROJECT/name,output/name)
    for name in ['assets/export-template.xlsx','vendor/fflate.js']:
        destination=output/name;destination.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(PROJECT/name,destination)
def build(config,force=False):
    PRIVATE.mkdir(exist_ok=True)
    selected=discover(config);manifest=[];docs={}
    for role,paths in selected.items():
        docs[role]=[]
        for path in paths:
            sha=digest(path);manifest.append({'role':role,'name':path.name,'period':period(path),'sha256':sha,'size':path.stat().st_size,'modified':datetime.fromtimestamp(path.stat().st_mtime,timezone(timedelta(hours=-6))).isoformat()})
            docs[role].append((path,sha))
    fingerprint=hashlib.sha256(json.dumps({'sources':manifest,'version':VERSION,'core':digest(PROJECT/'core.js'),'pipeline':digest(Path(__file__))},sort_keys=True).encode()).hexdigest()
    output=Path(config['output']);current=output/'data/version.json'
    if not force and current.exists() and json.loads(current.read_text(encoding='utf-8')).get('id')==fingerprint:
        sync_site(output)
        print('Sin cambios en las fuentes. Se conserva el tablero.');return {'changed':False,'id':fingerprint}
    print('Leyendo fuentes modificadas...',flush=True)
    loaded={role:[read_cached(p,role,sha) for p,sha in items] for role,items in docs.items()}
    standards=loaded['bpi']+loaded['records']+loaded['agency']
    # Core build remains the same tested rejection logic used by browser and Excel.
    incoming=PRIVATE/'entrada.json';write_json(incoming,standards)
    node=shutil.which('node') or str(Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe')
    result=subprocess.run([node,str(PROJECT/'automation/build-model.cjs'),str(incoming)],capture_output=True,text=True,encoding='utf-8',timeout=600)
    if result.returncode: raise ValueError(result.stderr[-4000:])
    model=json.loads(result.stdout);model=enrich(model,loaded['stores'][0],loaded['attendance'],loaded['bpi'])
    count=len(model['rows'])
    if len({r['folio'] for r in model['rows']})!=count: raise ValueError('El cruce duplicó folios. No se reemplazaron los datos.')
    for role in ('stores','attendance'):
        model['files'] += [{'name':d['name'],'role':role,'sheet':d['sheet'],'rows':len(d['rows'])} for d in loaded[role]]
    months={m:{'base':sum(r['month']==m for r in model['rows']),'cases':sum(r['month']==m and r['reject'] for r in model['rows']),'records':sum(r['month']==m and r['hasRecords'] for r in model['rows']),'lastBpiDate':max(r['created'] for r in model['rows'] if r['month']==m)} for m in model['months']}
    for item in manifest:
        if item['role'] in ('bpi','records'):
            source=next((d for d in loaded[item['role']] if d['name']==item['name']),None)
            dates=[iso(r.get('created')) for r in source['rows']] if source else []
            item['lastDate']=max((d for d in dates if d),default='')
    model['automation']={'generated':now().isoformat(),'id':fingerprint,'sources':manifest,'months':months,'sourcePolicy':'Una instantánea más reciente por mes; base y consolidado únicos; no se modifican los originales.'}
    # Check signatures again before exposing the new generation.
    for role,items in docs.items():
        for p,sha in items:
            if digest(p)!=sha: raise ValueError(f'{p.name}: cambió durante el proceso. Se conserva la versión anterior.')
    output.mkdir(parents=True,exist_ok=True)
    sync_site(output)
    # Immutable generation first, version pointer last: readers never see partial JSON.
    name=f'model-{fingerprint[:20]}.json';write_json(output/'data'/name,model)
    version={'id':fingerprint,'model':name,'generated':model['automation']['generated'],'months':months}
    write_json(current,version);write_json(PRIVATE/'ultima_actualizacion.json',{'ok':True,**version,'enrichment':model['enrichment'],'sources':manifest})
    print(json.dumps({'changed':True,'months':months,'attribution':model['enrichment']['totals'],'output':str(output)},ensure_ascii=False),flush=True)
    return {'changed':True,**version}

def run(config,force=False):
    PRIVATE.mkdir(exist_ok=True);lock=PRIVATE/'update.lock'
    try:
        # Windows advisory byte lock is released by the OS after a crash.
        import msvcrt
        with lock.open('a+b') as guard:
            if guard.seek(0,2)==0: guard.write(b'0');guard.flush()
            guard.seek(0)
            try: msvcrt.locking(guard.fileno(),msvcrt.LK_NBLCK,1)
            except OSError: print('Ya hay una actualización en curso.');return
            try:
                result=build(config,force)
                write_json(Path(config['output'])/'data/status.json',{'ok':True,'checked':now().isoformat()})
                return result
            finally: guard.seek(0);msvcrt.locking(guard.fileno(),msvcrt.LK_UNLCK,1)
    except Exception as exc:
        error={'ok':False,'checked':now().isoformat(),'error':str(exc)}
        write_json(PRIVATE/'ultimo_error.json',error)
        write_json(Path(config['output'])/'data/status.json',error)
        raise

def main():
    p=argparse.ArgumentParser();p.add_argument('--config',type=Path,default=CONFIG);p.add_argument('--force',action='store_true');p.add_argument('--watch',action='store_true');args=p.parse_args()
    config=json.loads(args.config.read_text(encoding='utf-8'))
    while True:
        try: run(config,args.force);args.force=False
        except Exception as exc:
            print('No se actualizó: '+str(exc),file=sys.stderr,flush=True)
            if not args.watch: return 1
        if not args.watch: return 0
        time.sleep(max(30,int(config.get('intervalSeconds',300))))
if __name__=='__main__': sys.exit(main())
